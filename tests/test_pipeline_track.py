"""Track boundaries use synthetic documents and never call a paid model."""
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import desktop_tasks
from pipeline_track import result_track, resolve_run_track, track_contract
from result_package import initialize_result_package, resolve_analysis_root


class PipelineTrackTests(unittest.TestCase):
    def test_explicit_track_is_independent_of_environment(self):
        with patch.dict(os.environ, {"RATOMIZER_FUNCTIONAL_EXTRACT": "0"}):
            self.assertEqual(resolve_run_track("functional", skip_review=True), "functional")
        with patch.dict(os.environ, {"RATOMIZER_FUNCTIONAL_EXTRACT": "1"}):
            self.assertEqual(resolve_run_track("legacy_a", skip_review=True), "legacy_a")
        with self.assertRaisesRegex(ValueError, "逐原子审查"):
            resolve_run_track("functional", skip_review=False)

    def test_functional_run_reuses_parser_then_explicit_legacy_run_regenerates(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "input.docx"
            doc = Document()
            doc.add_heading("4 Event log", 1)
            doc.add_paragraph("The meter shall retain 10 events after a power failure.")
            doc.save(source)
            out = root / "out"
            options = dict(skip_review=True, kb_paths=[], llm_route="stub")
            with patch.dict(os.environ, {"RATOMIZER_FUNCTIONAL_EXTRACT": "1"}), \
                    patch("desktop_tasks._stage_summary", return_value={}):
                first = desktop_tasks.run_pipeline_task(source, out, track="functional", **options)
                self.assertEqual(first["track"], "functional")
                self.assertIsNone(first["review"])
                with patch("desktop_tasks.run_atomizer_pipeline", side_effect=AssertionError("parser reran")):
                    second = desktop_tasks.run_pipeline_task(source, out, track="functional", **options)
                self.assertEqual(second["manifest"]["resume_action"], "skipped")
                # Empty tables are valid for prose; a missing file still invalidates reuse.
                (out / "table_items.jsonl").unlink()
                with patch("desktop_tasks.run_atomizer_pipeline", wraps=desktop_tasks.run_atomizer_pipeline) as parse:
                    desktop_tasks.run_pipeline_task(source, out, track="functional", **options)
                    parse.assert_called_once()
                self.assertFalse((out / "atomic_requirements.jsonl").exists())
                self.assertFalse((out / "llm_tasks.jsonl").exists())
                legacy = desktop_tasks.run_pipeline_task(source, out, track="legacy_a", **options)
                self.assertEqual(legacy["manifest"]["track"], "legacy_a")
                self.assertTrue((out / "atomic_requirements.jsonl").exists())
                self.assertTrue((out / "llm_tasks.jsonl").exists())
                # Switching back must not leave old atomic inputs for readers.
                desktop_tasks.run_pipeline_task(source, out, track="functional", **options)
                self.assertFalse((out / "atomic_requirements.jsonl").exists())
                self.assertFalse((out / "llm_tasks.jsonl").exists())

    def test_legacy_entrypoints_refuse_functional_before_work(self):
        from assemble_spec import assemble
        from engineering_composer import compose_engineering_requirements
        from llm_pipeline import run_review_pipeline
        from ai_extract import run_ai_extract
        from functional_synthesis import run_functional_synthesis

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            (out / "manifest.json").write_text('{"track":"functional"}')
            calls = [
                lambda: assemble(out, None, source="test", extracted_at="test"),
                lambda: compose_engineering_requirements(out),
                lambda: run_review_pipeline(out, route="stub"),
                lambda: run_ai_extract(out, route="stub"),
                lambda: run_functional_synthesis(out, route="stub"),
            ]
            for call in calls:
                with self.subTest(call=call), self.assertRaisesRegex(ValueError, "完整需求流程"):
                    call()
            with patch("desktop_tasks.functional_extract_task") as extract:
                with self.assertRaisesRegex(ValueError, "assemble"):
                    desktop_tasks.chain_task(out, stages=["functional-extract", "assemble"])
                extract.assert_not_called()
            with patch.dict(os.environ, {"RATOMIZER_FUNCTIONAL_EXTRACT": "0"}):
                with self.assertRaisesRegex(RuntimeError, "ai-extract"):
                    desktop_tasks.chain_task(out, stages=["ai-extract"])

    def test_explicit_legacy_and_undeclared_results_remain_usable(self):
        from pipeline_track import require_legacy_track

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            for manifest in ({}, {"track": "legacy_a"}):
                (out / "manifest.json").write_text(json.dumps(manifest))
                require_legacy_track(out, "assemble")
                self.assertTrue(track_contract(out)["legacy_stages_allowed"])
            (out / "manifest.json").write_text('{"track":"invalid"}')
            with self.assertRaisesRegex(ValueError, "未知"):
                result_track(out)

    def test_package_root_and_analysis_root_report_same_track(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            source.write_bytes(b"synthetic")
            package = root / "package"
            initialize_result_package(package, input_path=source, requested_stages=["atomize"])
            analysis = resolve_analysis_root(package)
            (analysis / "manifest.json").write_text('{"track":"functional"}')
            self.assertEqual(track_contract(package), track_contract(analysis))
            self.assertEqual(result_track(package), "functional")

    def test_api_describes_functional_and_historical_results(self):
        from tests.test_api_server import _claim_api, _http_json

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            with _claim_api(out) as url:
                status, old = _http_json(url, "/pipeline-track")
                self.assertEqual(status, 200)
                self.assertEqual(old["track"], "unknown")
                self.assertEqual(old["source"], "undeclared")
                (out / "manifest.json").write_text('{"track":"functional"}')
                status, current = _http_json(url, "/pipeline-track")
                self.assertEqual(status, 200)
                self.assertFalse(current["legacy_stages_allowed"])
                self.assertEqual(current["functional_requirements_endpoint"], "/functional-requirements")
                # The existing legacy list contract stays readable.
                status, rows = _http_json(url, "/requirements")
                self.assertEqual((status, rows), (200, []))
