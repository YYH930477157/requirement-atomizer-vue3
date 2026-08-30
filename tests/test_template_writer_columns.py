"""模板写入器列位解析（template_writer/v2）：需求 sheet 布局方差修复。

门禁 FAIL 根因（docs/ws0-gate-result-2026-08-30.md）：事件需求/状态字需求 sheet
是 8 列布局（无「需求模版」列），写入器固定列位（_COL_ANSWER=6）把正文写进
「说明」列、「需求」列空着——门禁读 58+2=60 条空正文行判 FAIL。

v2 语义：每个需求 sheet 先按**表头名**解析列位（序号/子模块/描述/需求/说明/
是否客户需求/客户需求章节/驱动硬件；需求模版列在场与否不再影响列位）；表头
解析失败（计量需求拆分列等别名失效场景）回退固定列契约（v1 行为钉住不变）；
全空条目（描述/正文/说明皆空）跳过并审计计数，不再产生纯序号空行。
"""
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import template_writer as tw

# 标准 10 列布局（含需求模版）
HEADER_STD = ["关闭", "序号", "子模块", "描述", "需求模版", "需求",
              "说明、示例、注意事项", "是否客户需求", "客户需求章节", "驱动/硬件相关"]
# 8 列布局（无需求模版——事件需求/状态字需求实测形态）
HEADER_NARROW = ["关闭", "序号", "子模块", "描述", "需求",
                 "说明、示例、注意事项", "是否客户需求", "客户需求章节"]
# 计量需求型：表头被电表类型列拆分，无法按名解析（契约回退场景）
HEADER_SPLIT = ["关闭", "序号", "子模块", "描述", "1P2W_SP", "3P4W_DC", "3P4W_LVCT",
                "说明、示例、注意事项", "是否客户需求", "客户需求章节"]


def _item(module: str = "事件记录", description: str = "事件记录存储",
          requirement: str = "目标：实现事件记录存储。", objective: str = "实现事件记录存储") -> dict:
    return {"module": module, "submodule": module, "description": description,
            "requirement": requirement, "objective": objective,
            "source_section": "8.9"}


def _make_template(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "事件需求"
    ws.append(HEADER_NARROW)
    ws.append(["", "1", "事件", "事件log filter：", "支持", "建议都实现该功能", "", ""])
    ws2 = wb.create_sheet("系统需求")
    ws2.append(HEADER_STD)
    ws2.append(["", "1", "基本参数", "电表类型：", "1P2W", "1P2W", "说明", ""])
    ws4 = wb.create_sheet("时钟需求")
    ws4.append(HEADER_STD)
    ws4.append(["", "1", "时钟", "历法：", "公历", "公历", "", ""])
    ws3 = wb.create_sheet("计量需求")
    ws3.append(HEADER_SPLIT)
    ws3.append(["", "1", "基本参数", "火线采样类型：", "1", "1", "①CT采样", ""])
    wb.save(path)


class ColumnResolutionTests(unittest.TestCase):
    def test_narrow_sheet_body_lands_in_requirement_column(self) -> None:
        """8 列 sheet：正文必须落在「需求」列（v1 会错位写进「说明」列）。"""
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "tpl.xlsx"
            out = Path(tmp) / "written.xlsx"
            _make_template(src)
            report = tw.append_analysis_to_template(
                src, [_item()], out)
            ws = load_workbook(out)["事件需求"]
            rows = list(ws.iter_rows(values_only=True))
            appended = rows[-1]
            self.assertEqual(str(appended[4]), "目标：实现事件记录存储。", "正文应在第5列「需求」")
            self.assertEqual(str(appended[5]), "功能目标：实现事件记录存储", "说明应在第6列")
            self.assertEqual(report["appended_by_sheet"].get("事件需求"), 1)
            self.assertEqual(
                (report.get("column_resolution") or {}).get("事件需求"), "header")

    def test_std_sheet_positions_unchanged(self) -> None:
        """标准 10 列 sheet（含需求模版列）：表头解析列位与 v1 固定列位一致
        （描述4/正文6/说明7）——表头解析对标准布局零漂移。"""
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "tpl.xlsx"
            out = Path(tmp) / "written.xlsx"
            _make_template(src)
            tw.append_analysis_to_template(
                src, [_item(module="时钟")], out)
            ws = load_workbook(out)["时钟需求"]
            appended = list(ws.iter_rows(values_only=True))[-1]
            self.assertEqual(str(appended[3]), "事件记录存储")
            self.assertEqual(str(appended[5]), "目标：实现事件记录存储。")
            self.assertEqual(str(appended[6]), "功能目标：实现事件记录存储")

    def test_split_header_sheet_falls_back_to_contract(self) -> None:
        """表头不可解析（计量需求拆分列）：回退固定列契约——v1 行为钉住。"""
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "tpl.xlsx"
            out = Path(tmp) / "written.xlsx"
            _make_template(src)
            report = tw.append_analysis_to_template(
                src, [_item(module="计量")], out)
            ws = load_workbook(out)["计量需求"]
            appended = list(ws.iter_rows(values_only=True))[-1]
            # 契约列位：正文=6（写入器与读取器共用同一权威）
            self.assertEqual(str(appended[5]), "目标：实现事件记录存储。")
            self.assertEqual(
                (report.get("column_resolution") or {}).get("计量需求"), "contract")

    def test_fully_empty_item_skipped_with_audit(self) -> None:
        """全空条目（无描述/正文/说明）不得产生纯序号空行——跳过并审计。"""
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "tpl.xlsx"
            out = Path(tmp) / "written.xlsx"
            _make_template(src)
            empty = {"module": "事件记录", "submodule": "", "description": "",
                     "requirement": "", "analysis_notes": ""}
            report = tw.append_analysis_to_template(
                src, [empty, _item()], out)
            self.assertEqual(report.get("skipped_empty_items"), 1)
            self.assertEqual(report["appended_by_sheet"].get("事件需求"), 1)


if __name__ == "__main__":
    unittest.main()
