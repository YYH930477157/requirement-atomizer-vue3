"""待核成文（partial export）全链：守恒未闭合/直抽 partial 时分析·成文照跑并标 partial。

方案 docs/pending-export-plan-2026-08-31.md v2。语义：
- FunctionalConservationError 与 execution_status=partial（mixed）在开关开（默认）时
  走旁路：分析/成文/澄清跑完、manifest 记 partial（非 ok——不可当干净复用）、
  chain 载荷 conservation_blocked（保留未闭合义）+ partial_export + 待核行数；
- execution_status=failed 仍整段拦（旁路不放行数据不完整）；
- RATOMIZER_PARTIAL_EXPORT=0 回到今日行为（跳过 + failed）；
- xlsx 行上带「⚠待核（失败类）」，干净行无标记；draft:true 不打。
"""
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import desktop_tasks
from openpyxl import load_workbook

HEADER = ["关闭", "序号", "子模块", "描述", "需求", "说明、示例、注意事项",
          "是否客户需求", "客户需求章节"]


def _make_template(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "事件需求"
    ws.append(HEADER)
    ws.append(["", "1", "事件", "事件log filter：", "支持", "建议都实现", "", ""])
    wb.save(path)


def _seed_fr(out: Path, *, execution_status: str = "ok",
             conservation_ok: bool = False, route: str = "stub",
             stub_f2: bool = False) -> None:
    (out / "blocks.jsonl").write_text(
        '{"block_id":"B1","type":"paragraph","section_path":["4.1"],'
        '"text":"The meter shall log events.","order":1}\n', encoding="utf-8")
    (out / "chunks.jsonl").write_text(
        '{"section_path":["4.1"],"heading":"4.1",'
        '"text":"The meter shall log events.","block_ids":["B1"]}\n', encoding="utf-8")
    (out / "table_items.jsonl").write_text("", encoding="utf-8")
    import functional_extract as fe

    section = {"section_id": "4.1", "section_path": ["4.1"], "heading": "4.1",
               "text": "The meter shall log events.", "block_ids": ["B1"]}
    f2 = {"functional_requirement_id": "F2", "module": "事件记录", "submodule": "事件记录",
          "title": "上报事件", "description": "上报事件",
          "requirement": "目标：实现事件上报。",
          "objective": "实现事件上报",
          "source_block_ids": ["B1"], "source_quote": "The meter shall log events.",
          "source_section": "4.1", "type": "functional", "priority": "P1"}
    if stub_f2:
        # v2：mixed 载荷中包级 stub 退化的占位条目——与 _stub_item 同形状
        objective, behaviors = fe._stub_shape(section)
        f2.update({"objective": objective, "behaviors": behaviors,
                   "title": "4.1", "description": objective})
    items = [
        {"functional_requirement_id": "F1", "module": "事件记录", "submodule": "事件记录",
         "title": "记录事件", "description": "记录事件",
         "requirement": "目标：实现事件记录。",
         "objective": "实现事件记录",
         "source_block_ids": ["B1"], "source_quote": "The meter shall log events.",
         "source_section": "4.1", "type": "functional", "priority": "P1"},
        f2,
    ]
    fr = {
        "producer": "functional-extract/v1+test", "route": route,
        "route_requested": route, "execution_status": execution_status,
        "schema_version": "functional-requirements/v1", "items": items,
        "conservation": {
            "ok": conservation_ok,
            "checks": {
                "clause_coverage": {"ok": True, "uncovered_sections": []},
                "obligation_coverage": {"ok": True, "uncovered_obligations": []},
                "evidence_presence": {
                    "ok": conservation_ok,
                    "binding_mismatches": ([] if conservation_ok else [{
                        "functional_requirement_id": "F1",
                        "reason": "narrative_covers_other_clauses_not_declared",
                    }]),
                    "evidence_mismatches": [], "items_without_evidence": [],
                },
                "duplicates": {"ok": True, "groups": []},
                "preservation": {"ok": True, "blocking_losses": [], "warning_losses": []},
            },
        },
    }
    (out / "functional_requirements.json").write_text(
        json.dumps(fr, ensure_ascii=False), encoding="utf-8")


class PartialExportChainTests(unittest.TestCase):
    def _run(self, out: Path, tpl: Path, *, env: str | None = "1"):
        patches = []
        if env is not None:
            patches.append(mock.patch.dict("os.environ", {"RATOMIZER_PARTIAL_EXPORT": env}))
        with patches[0] if patches else _null():
            return desktop_tasks.chain_task(
                out, stages=["requirements-analysis", "template-write"],
                route="openai_compatible", template_path=tpl)

    def test_unclosed_conservation_runs_stages_partial_and_marks_rows(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr(out, execution_status="ok", conservation_ok=False)
            payload = self._run(out, tpl)

            self.assertTrue(payload.get("partial_export"), "旁路应开")
            self.assertTrue(payload.get("conservation_blocked"), "未闭合义保留")
            manifest = desktop_tasks.read_run_manifest(out)["stages"]
            self.assertEqual(manifest["requirements-analysis"]["status"], "partial")
            self.assertEqual(manifest["template-write"]["status"], "partial")
            xlsx = out / "软件需求列表-成文.xlsx"
            self.assertTrue(xlsx.exists(), "成文应写出")
            rows = list(load_workbook(xlsx)["事件需求"].iter_rows(values_only=True))
            texts = [str(r[5] or "") for r in rows]  # 说明列
            self.assertTrue(any("待核" in t and "绑定" in t for t in texts),
                            "F1 行说明列应带待核+绑定失配")
            self.assertTrue(any(t and "待核" not in t for t in texts[2:]),
                            "F2 干净行不应带待核标记")
            self.assertFalse(any("draft" in t for t in texts), "不打 draft 水印")
            self.assertGreaterEqual(payload.get("pending_marked_rows", 0), 1)
            # 2026-09-05 P2：链尾产物单源信号——守恒未闭合必须在链载荷可见
            # （GUI 日常链无分析/成文阶段时这是运行页唯一信号）
            self.assertIn("功能需求守恒核对未闭合",
                          str(payload.get("functional_conservation_error")))

    def test_extraction_partial_mixed_also_bypasses(self) -> None:
        """grok 第 3 条：execution_status=partial（mixed）走旁路——SBD 主形态。"""
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr(out, execution_status="partial")
            payload = self._run(out, tpl)
            self.assertTrue(payload.get("partial_export"))
            self.assertTrue((out / "软件需求列表-成文.xlsx").exists())

    def test_mixed_stub_rows_marked_extract_degraded_even_when_conservation_ok(self) -> None:
        """v2：mixed 载荷的 stub 占位行挂 extract_degraded——与守恒状态独立。

        守恒闭合（ok=True）+ mixed：阶段照跑不记 partial（无守恒缺口），但降级行
        必须在 xlsx 上可见、行数进链载荷。
        """
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr(out, execution_status="partial", conservation_ok=True,
                     route="mixed", stub_f2=True)
            payload = self._run(out, tpl)
            self.assertFalse(payload.get("conservation_blocked"), "守恒闭合不应记未闭合")
            self.assertFalse(payload.get("functional_conservation_error"),
                             "守恒闭合不应有未闭合链层信号")
            self.assertEqual(payload.get("pending_marked_rows"), 1, "降级行数上报")
            manifest = desktop_tasks.read_run_manifest(out)["stages"]
            self.assertEqual(manifest["requirements-analysis"]["status"], "ok",
                             "守恒闭合的 mixed 不记 partial（降级不是守恒缺口）")
            xlsx = out / "软件需求列表-成文.xlsx"
            rows = list(load_workbook(xlsx)["事件需求"].iter_rows(values_only=True))
            texts = [str(r[5] or "") for r in rows]
            self.assertTrue(any("待核" in t and "抽取降级" in t for t in texts),
                            "stub 占位行说明列应带「⚠待核（抽取降级）」")
            self.assertIn(
                "守恒待核", load_workbook(xlsx).sheetnames,
                "extract_degraded-only 也必须建「守恒待核」清单",
            )

    def test_mixed_stub_row_merges_with_conservation_class(self) -> None:
        """v2：同一行既守恒待核又抽取降级——两类去重并存。"""
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            # F1 = binding 失配点名；F2 = stub 占位 → 两行各一类；再让 F2 也被
            # binding 点名构造同类合并场景：把 binding 指向 F2
            _seed_fr(out, execution_status="partial", conservation_ok=False,
                     route="mixed", stub_f2=True)
            # 把 binding_mismatches 的点名从 F1 改为 F2（占位行叠加守恒类）
            fr_path = out / "functional_requirements.json"
            fr = json.loads(fr_path.read_text(encoding="utf-8"))
            fr["conservation"]["checks"]["evidence_presence"]["binding_mismatches"] = [{
                "functional_requirement_id": "F2", "reason": "r"}]
            fr_path.write_text(json.dumps(fr, ensure_ascii=False), encoding="utf-8")
            payload = self._run(out, tpl)
            self.assertTrue(payload.get("partial_export"))
            xlsx = out / "软件需求列表-成文.xlsx"
            rows = list(load_workbook(xlsx)["事件需求"].iter_rows(values_only=True))
            texts = [str(r[5] or "") for r in rows]
            merged = [t for t in texts if "抽取降级" in t]
            self.assertTrue(merged, "F2 行应有抽取降级类")
            self.assertIn("绑定失配", merged[0], "两类并存同一行去重展示")

    def test_gap_only_unclosed_template_write_self_reports_partial(self) -> None:
        """2026-09-05 P3：gap-only 边缘——守恒未闭合但失败面全是零 FRE 缺口
        （marked_rows=0、只有「守恒待核」清单的 gap 行）。

        template-write 不进 GUI 默认链，此形态走 CLI/门禁——成文也必须自报
        unclosed_basis（链记 partial），否则 manifest 记 ok、与清单 sheet 的
        红字总述矛盾。分析不进链：engineering_analysis 直接落盘，隔离成文自报。
        """
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr(out, execution_status="ok", conservation_ok=False)
            # 失败面改成纯缺口：evidence 恢复 ok，唯一失败 = 条款覆盖缺口章节
            fr_path = out / "functional_requirements.json"
            fr = json.loads(fr_path.read_text(encoding="utf-8"))
            checks = fr["conservation"]["checks"]
            checks["evidence_presence"]["ok"] = True
            checks["evidence_presence"]["binding_mismatches"] = []
            checks["clause_coverage"]["ok"] = False
            checks["clause_coverage"]["uncovered_sections"] = [
                {"section_id": "9.9", "heading": "9.9 Spare"}]
            fr_path.write_text(json.dumps(fr, ensure_ascii=False), encoding="utf-8")
            (out / "engineering_analysis.json").write_text(
                json.dumps({"items": fr["items"]}, ensure_ascii=False),
                encoding="utf-8")
            with mock.patch.dict("os.environ", {"RATOMIZER_PARTIAL_EXPORT": "1"}):
                payload = desktop_tasks.chain_task(
                    out, stages=["template-write"],
                    route="openai_compatible", template_path=tpl)
            self.assertTrue(payload.get("conservation_blocked"),
                            "gap-only 也要自报未闭合（清单 sheet 红字在、链不能记 ok）")
            self.assertTrue(payload.get("partial_export"))
            manifest = desktop_tasks.read_run_manifest(out)["stages"]
            self.assertEqual(manifest["template-write"]["status"], "partial")
            report = payload["results"]["template-write"]["report"]
            block = report["conservation_pending_export"]
            self.assertEqual(block["marked_rows"], 0, "失败面没有可挂 FRE 的行")
            self.assertEqual(block["gap_rows"], 1, "缺口行进清单")
            self.assertGreaterEqual(int(block["pending_sheet_rows"] or 0), 1)
            # 链尾产物单源信号在本形态同样可见（run 页提示数据源）
            self.assertIn("功能需求守恒核对未闭合",
                          str(payload.get("functional_conservation_error")))

    def test_extraction_failed_still_blocks(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr(out, execution_status="failed")
            payload = self._run(out, tpl)
            self.assertFalse(payload.get("partial_export"))
            self.assertTrue(payload.get("conservation_blocked"))
            self.assertFalse((out / "软件需求列表-成文.xlsx").exists())
            manifest = desktop_tasks.read_run_manifest(out)["stages"]
            self.assertEqual(manifest["requirements-analysis"]["status"], "failed")

    def test_switch_off_restores_today_behavior(self) -> None:
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            tpl = out / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr(out, execution_status="ok", conservation_ok=False)
            payload = self._run(out, tpl, env="0")
            self.assertFalse(payload.get("partial_export"))
            self.assertTrue(payload.get("conservation_blocked"))
            self.assertFalse((out / "软件需求列表-成文.xlsx").exists())
            manifest = desktop_tasks.read_run_manifest(out)["stages"]
            self.assertEqual(manifest["requirements-analysis"]["status"], "failed")
            self.assertEqual(manifest["template-write"]["status"], "failed")
            self.assertTrue(payload["results"]["template-write"].get(
                "skipped_due_to_conservation"))


class _null:
    def __enter__(self):
        return None

    def __exit__(self, *a):
        return False


class ProducerStampTests(unittest.TestCase):
    def test_producer_carries_version_and_env_value(self) -> None:
        """grok 第 2 条：标记算法身份与开关有效值进两阶段 producer（=0/=1 隔离）。"""
        from functional_extract import CONSERVATION_PARTIAL_EXPORT_VERSION

        with mock.patch.dict("os.environ", {"RATOMIZER_PARTIAL_EXPORT": "1"}):
            on = desktop_tasks.stage_producer("requirements-analysis")
            on_tw = desktop_tasks.stage_producer("template-write")
        with mock.patch.dict("os.environ", {"RATOMIZER_PARTIAL_EXPORT": "0"}):
            off = desktop_tasks.stage_producer("requirements-analysis")
        # 两态都带版本+有效值（靠值隔离，=0/=1 不互复用——producer 在
        # _stage_is_reusable 里逐字节比较，desktop_tasks.py:1707）
        self.assertIn(CONSERVATION_PARTIAL_EXPORT_VERSION, on)
        self.assertIn("partial-export-True", on)
        self.assertIn(CONSERVATION_PARTIAL_EXPORT_VERSION, on_tw)
        self.assertIn("partial-export-False", off)
        self.assertNotEqual(on, off)

    def test_partial_stage_not_reusable_as_clean(self) -> None:
        """partial 的成文不得被 stage_is_reusable 当干净代复用。"""
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            desktop_tasks.update_run_manifest(
                out, "template-write", "partial", action="ran")
            self.assertFalse(desktop_tasks.stage_is_reusable(out, "template-write"))


# ---------------------------------------------------------------------------
# v3 Task 1：package_v1 寻址——桌面新跑的 FR 在 .ratomizer/pipeline/，裸根读漏标记
# ---------------------------------------------------------------------------

def _init_package(root: Path) -> Path:
    from result_package import initialize_result_package

    src = root / "standard.docx"
    src.write_bytes(b"docx-fixture")
    initialize_result_package(
        root, input_path=src,
        requested_stages=["requirements-analysis", "template-write"],
    )
    return root


def _seed_fr_governed(root: Path, **kwargs) -> None:
    from result_package import governed_artifact_path

    _seed_fr(root, **kwargs)  # 先写根目录（沿用既有夹具的 blocks/chunks 等）
    # 桌面 package_v1：pipeline 产物不在根目录。FR 必须 unlink 才能逼出裸读漏洞；
    # chunks/blocks 一并搬进 pipeline，让 load_clauses / gaps 走真实 governed 路径。
    for name in ("functional_requirements.json", "blocks.jsonl",
                 "chunks.jsonl", "table_items.jsonl"):
        src = root / name
        if not src.is_file():
            continue
        dest = governed_artifact_path(root, name, category="pipeline", for_write=True)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
        src.unlink()


class PartialExportPackageV1Tests(unittest.TestCase):
    def test_package_v1_unclosed_marks_rows_and_sets_partial_export(self) -> None:
        """package_v1：根目录无 FR、pipeline 里有未闭合产物 → 成文必须带 ⚠待核。

        v2 缺陷：attach / unclosed_basis 裸读根目录 → package_v1 桌面跑成文
        照出、行不标、UI 当干净表。修复后 governed 双路径读取。
        """
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            _init_package(root)
            tpl = root / "tpl.xlsx"
            _make_template(tpl)
            _seed_fr_governed(root, execution_status="ok", conservation_ok=False)
            # 无原子：根目录也不得有 functional_requirements.json
            self.assertFalse((root / "functional_requirements.json").exists())
            payload = PartialExportChainTests()._run(root, tpl)
            self.assertTrue(payload.get("partial_export"))
            self.assertGreater(int(payload.get("pending_marked_rows") or 0), 0)
            xlsx = root / "软件需求列表-成文.xlsx"
            self.assertTrue(xlsx.is_file())
            wb = load_workbook(xlsx)
            self.assertIn("守恒待核", wb.sheetnames, "package_v1 成文必须含清单 sheet")
            notes = []
            for ws in wb.worksheets:
                if ws.title == "守恒待核":
                    continue
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row and any(c not in (None, "") for c in row):
                        notes.append(" ".join(str(c or "") for c in row))
            self.assertTrue(
                any("⚠待核" in text for text in notes),
                "package_v1 成文必须带待核前缀，不能是干净表",
            )


if __name__ == "__main__":
    unittest.main()
